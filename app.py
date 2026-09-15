import streamlit as st
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from io import BytesIO
import re
import os
import json
import urllib.request
import urllib.error


# ============================================================
# MATERIAL SCANNER AGENT
# ============================================================
# AI-powered Material Pre-Review Prototype for AZ All Hackathon
#
# Current stage:
#   1. PowerPoint extraction
#   2. Rule-based keyword screening
#   3. Compliance pre-review
#   4. Evidence / Context Extraction
#   5. AI Context Analysis
#
# IMPORTANT:
# This prototype does NOT make a final compliance decision.
# AI analysis is contextual analysis only.
# Human review remains required for applicable content.
# ============================================================


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="Material Scanner Agent",
    page_icon="🔎",
    layout="wide"
)


# ============================================================
# APPLICATION INFORMATION
# ============================================================

APP_NAME = "Material Scanner Agent"

APP_SUBTITLE = (
    "AI-powered Material Pre-Review Prototype for AZ All Hackathon"
)


# ============================================================
# IT / SUPPORT KEYWORDS
# ============================================================

IT_KEYWORDS = [
    "it support",
    "service desk",
    "self service",
    "askaz",
    "site it",
    "it information",
    "support",
    "uniflow",
    "printer",
    "printing",
    "meeting room",
    "myaz",
    "sspr",
]


# ============================================================
# PROMOTIONAL KEYWORDS
# ============================================================

PROMOTIONAL_KEYWORDS = [
    "promotion",
    "promotional",
    "product offer",
    "special offer",
    "discount",
    "campaign",
    "price",
    "pricing",
    "buy now",
    "limited offer",
]


# ============================================================
# COMPLIANCE-SENSITIVE SCREENING TERMS
# ============================================================
#
# These are indicators only.
# They are NOT compliance rules.
# ============================================================

COMPLIANCE_REVIEW_TERMS = [
    "product",
    "indication",
    "efficacy",
    "safety",
    "clinical",
    "dose",
    "dosage",
    "treatment",
    "benefit",
    "approved",
    "patient",
    "prescribing",
    "prescription",
    "adverse event",
    "contraindication",
]


# ============================================================
# TEXT NORMALIZATION
# ============================================================

def normalize_text(text):
    if not text:
        return ""

    text = str(text)
    text = text.replace("\xa0", " ")
    text = re.sub(r"\s+", " ", text)

    return text.strip()


# ============================================================
# POWERPOINT SHAPE TEXT EXTRACTION
# ============================================================

def extract_text_from_shape(shape):
    """
    Recursively extract text from:
    - text boxes
    - group shapes
    - tables
    """

    texts = []

    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        for child_shape in shape.shapes:
            texts.extend(
                extract_text_from_shape(child_shape)
            )
        return texts

    if getattr(shape, "has_text_frame", False):
        text = normalize_text(shape.text)

        if text:
            texts.append(text)

    if getattr(shape, "has_table", False):
        for row in shape.table.rows:
            for cell in row.cells:
                text = normalize_text(cell.text)

                if text:
                    texts.append(text)

    return texts


# ============================================================
# POWERPOINT EXTRACTION
# ============================================================

def extract_slides(pptx_file):
    """
    Extract text slide-by-slide from a PowerPoint PPTX.
    """

    pptx_file.seek(0)

    presentation = Presentation(pptx_file)
    slides = []

    for slide_number, slide in enumerate(
        presentation.slides,
        start=1
    ):
        slide_texts = []

        for shape in slide.shapes:
            slide_texts.extend(
                extract_text_from_shape(shape)
            )

        unique_texts = list(
            dict.fromkeys(slide_texts)
        )

        slides.append(
            {
                "slide": slide_number,
                "texts": unique_texts
            }
        )

    return slides


# ============================================================
# COMBINE ALL TEXT
# ============================================================

def combine_all_text(slides):
    all_text = []

    for slide in slides:
        all_text.extend(
            slide.get("texts", [])
        )

    return " ".join(all_text)


# ============================================================
# KEYWORD MATCHING
# ============================================================

def find_keyword_matches(slides, keywords):
    """
    Search every slide for every keyword.

    Output is always sorted:
        1. Slide number
        2. Keyword
    """

    detected_keywords = []
    slide_references = []

    for slide in slides:
        slide_number = slide.get("slide")

        slide_text = " ".join(
            slide.get("texts", [])
        ).lower()

        for keyword in keywords:
            if keyword.lower() in slide_text:
                detected_keywords.append(keyword)

                slide_references.append(
                    (
                        slide_number,
                        keyword
                    )
                )

    detected_keywords = list(
        dict.fromkeys(detected_keywords)
    )

    slide_references = list(
        dict.fromkeys(slide_references)
    )

    slide_references.sort(
        key=lambda item: (
            int(item[0]),
            str(item[1]).lower()
        )
    )

    return (
        detected_keywords,
        slide_references
    )


# ============================================================
# MATERIAL ANALYSIS
# ============================================================

def analyze_material(slides):
    """
    Perform rule-based material screening.

    This function does NOT perform AI analysis.
    """

    combined_text = combine_all_text(
        slides
    ).lower()

    detected_it, it_slide_refs = (
        find_keyword_matches(
            slides,
            IT_KEYWORDS
        )
    )

    detected_promo, promo_slide_refs = (
        find_keyword_matches(
            slides,
            PROMOTIONAL_KEYWORDS
        )
    )

    detected_review, review_slide_refs = (
        find_keyword_matches(
            slides,
            COMPLIANCE_REVIEW_TERMS
        )
    )

    material_type = (
        "Promotional"
        if detected_promo
        else "Non-Promotional"
    )

    if detected_promo:
        risk_level = "Medium"
    elif detected_review:
        risk_level = "Review"
    else:
        risk_level = "Low"

    findings = []

    if detected_it:
        findings.append(
            "Material appears to contain internal IT/support information."
        )

    if detected_promo:
        findings.append(
            "Potential promotional content detected."
        )
    else:
        findings.append(
            "No obvious promotional keywords detected."
        )

    if detected_review:
        findings.append(
            "Content requiring further compliance interpretation was detected."
        )

    if not findings:
        findings.append(
            "No screening findings were generated."
        )

    return {
        "material_type": material_type,
        "risk_level": risk_level,
        "findings": findings,
        "it_keywords": detected_it,
        "promo_keywords": detected_promo,
        "review_keywords": detected_review,
        "it_slide_refs": it_slide_refs,
        "promo_slide_refs": promo_slide_refs,
        "review_slide_refs": review_slide_refs,
        "compliance_status": "Pending AI Review",
        "combined_text": combined_text,
    }


# ============================================================
# EVIDENCE / CONTEXT EXTRACTION
# ============================================================

def build_evidence_context(slides, analysis):
    """
    Build one evidence record per:
        Category + Slide

    Multiple keywords found on the same slide are grouped
    together so the AI receives the full slide context.
    """

    evidence_map = {}

    category_refs = [
        (
            "IT / Support",
            analysis.get("it_slide_refs", [])
        ),
        (
            "Promotional",
            analysis.get("promo_slide_refs", [])
        ),
        (
            "Compliance-Sensitive",
            analysis.get("review_slide_refs", [])
        ),
    ]

    for category, refs in category_refs:
        for slide_number, keyword in refs:

            key = (
                category,
                int(slide_number)
            )

            if key not in evidence_map:
                evidence_map[key] = {
                    "slide": int(slide_number),
                    "category": category,
                    "keywords": [],
                    "evidence": ""
                }

            if keyword not in evidence_map[key]["keywords"]:
                evidence_map[key]["keywords"].append(keyword)

    slides_by_number = {
        int(slide.get("slide")): slide
        for slide in slides
    }

    for item in evidence_map.values():

        slide = slides_by_number.get(
            item["slide"],
            {}
        )

        item["evidence"] = normalize_text(
            " ".join(
                slide.get("texts", [])
            )
        )

    evidence_items = list(
        evidence_map.values()
    )

    evidence_items.sort(
        key=lambda item: (
            int(item["slide"]),
            str(item["category"]).lower()
        )
    )

    return evidence_items


# ============================================================
# AI CONFIGURATION
# ============================================================
#
# Expected environment variables:
#
#   AI_API_URL
#   AI_API_KEY
#   AI_MODEL
#
# The endpoint is expected to support an
# OpenAI-compatible /chat/completions request.
#
# No key is hard-coded into this application.
# ============================================================

def get_ai_configuration():

    return {
        "api_url": os.getenv(
            "AI_API_URL",
            ""
        ).strip(),

        "api_key": os.getenv(
            "AI_API_KEY",
            ""
        ).strip(),

        "model": os.getenv(
            "AI_MODEL",
            ""
        ).strip(),
    }


# ============================================================
# AI CONTEXT ANALYSIS PROMPT
# ============================================================

def build_ai_context_prompt(
    evidence_items,
    analysis
):
    evidence_payload = []

    for item in evidence_items:
        evidence_payload.append(
            {
                "slide": item["slide"],
                "category": item["category"],
                "keywords": item["keywords"],
                "evidence": item["evidence"],
            }
        )

    evidence_json = json.dumps(
        evidence_payload,
        ensure_ascii=False,
        indent=2
    )

    analysis_json = json.dumps(
        {
            "material_type":
                analysis.get("material_type"),

            "risk_level":
                analysis.get("risk_level"),

            "findings":
                analysis.get("findings"),
        },
        ensure_ascii=False,
        indent=2
    )

    return f"""
You are the contextual analysis stage of Material Scanner Agent.

Analyze the supplied PowerPoint evidence and explain what the
detected content appears to mean in context.

IMPORTANT:
- Do NOT make a final compliance or regulatory decision.
- Do NOT invent facts.
- Keyword detection is only an indicator, not proof of non-compliance.
- Base the assessment only on the supplied slide evidence.
- If the evidence is insufficient, say "Insufficient evidence".
- Distinguish informational/internal content from promotional or
  potentially compliance-sensitive content.
- Keep the output concise and suitable for human review.

Rule-based analysis:
{analysis_json}

Evidence:
{evidence_json}

Return ONLY valid JSON in exactly this structure:

{{
  "overall_summary": "short evidence-based summary",
  "items": [
    {{
      "slide": 1,
      "category": "IT / Support",
      "keywords": ["keyword"],
      "context": "what the evidence appears to mean",
      "assessment": "Informational",
      "risk_level": "Low",
      "reason": "short evidence-based reason",
      "recommendation": "short recommended next step"
    }}
  ]
}}

Allowed assessment values:
- Informational
- Potentially Promotional
- Potentially Compliance-Sensitive
- Insufficient Evidence

Allowed risk_level values:
- Low
- Medium
- High
- Review
""".strip()


# ============================================================
# AI API CALL
# ============================================================

def call_ai_context_analysis(prompt, configuration):

    api_url = configuration["api_url"]
    api_key = configuration["api_key"]
    model = configuration["model"]

    if not api_url:
        raise RuntimeError(
            "AI_API_URL is not configured."
        )

    if not api_key:
        raise RuntimeError(
            "AI_API_KEY is not configured."
        )

    if not model:
        raise RuntimeError(
            "AI_MODEL is not configured."
        )

    request_body = {
        "model": model,
        "messages": [
            {
                "role": "system",
                "content": (
                    "You are a careful evidence-based "
                    "context analysis assistant."
                )
            },
            {
                "role": "user",
                "content": prompt
            }
        ],
        "temperature": 0.1,
    }

    body = json.dumps(
        request_body,
        ensure_ascii=False
    ).encode("utf-8")

    request = urllib.request.Request(
        api_url,
        data=body,
        method="POST"
    )

    request.add_header(
        "Content-Type",
        "application/json"
    )

    request.add_header(
        "Authorization",
        f"Bearer {api_key}"
    )

    try:

        with urllib.request.urlopen(
            request,
            timeout=60
        ) as response:

            response_text = response.read().decode(
                "utf-8"
            )

    except urllib.error.HTTPError as error:

        error_body = ""

        try:
            error_body = error.read().decode(
                "utf-8"
            )
        except Exception:
            pass

        raise RuntimeError(
            f"AI service returned HTTP {error.code}. "
            f"{error_body[:500]}"
        )

    except urllib.error.URLError as error:

        raise RuntimeError(
            f"Unable to connect to AI service: {error.reason}"
        )

    response_json = json.loads(
        response_text
    )

    try:
        content = (
            response_json
            ["choices"]
            [0]
            ["message"]
            ["content"]
        )
    except (
        KeyError,
        IndexError,
        TypeError
    ):
        raise RuntimeError(
            "AI response did not contain the expected "
            "chat-completions structure."
        )

    if isinstance(content, list):
        content = "".join(
            str(part.get("text", ""))
            if isinstance(part, dict)
            else str(part)
            for part in content
        )

    content = str(content).strip()

    content = re.sub(
        r"^```json\s*",
        "",
        content,
        flags=re.IGNORECASE
    )

    content = re.sub(
        r"^```\s*",
        "",
        content
    )

    content = re.sub(
        r"\s*```$",
        "",
        content
    )

    try:
        return json.loads(
            content.strip()
        )
    except json.JSONDecodeError:
        raise RuntimeError(
            "AI returned content that could not be parsed as JSON."
        )


# ============================================================
# AI CONTEXT ANALYSIS
# ============================================================

def ai_context_analysis(
    analysis,
    evidence_items
):
    """
    Run AI contextual analysis.

    The application continues to work when AI is not configured.
    """

    if not evidence_items:
        return {
            "status": "No Evidence",
            "overall_summary": (
                "No detected evidence was available for AI analysis."
            ),
            "items": []
        }

    configuration = get_ai_configuration()

    if not configuration["api_url"]:
        return {
            "status": "AI Not Configured",
            "overall_summary": (
                "AI Context Analysis is ready, but no approved "
                "AI endpoint has been configured in the runtime."
            ),
            "items": []
        }

    try:

        prompt = build_ai_context_prompt(
            evidence_items,
            analysis
        )

        result = call_ai_context_analysis(
            prompt,
            configuration
        )

        if not isinstance(result, dict):
            raise RuntimeError(
                "AI returned an unexpected JSON structure."
            )

        result.setdefault(
            "overall_summary",
            "AI analysis completed."
        )

        result.setdefault(
            "items",
            []
        )

        result["status"] = "Completed"

        return result

    except Exception as error:

        return {
            "status": "AI Error",
            "overall_summary": (
                "AI Context Analysis could not be completed."
            ),
            "items": [],
            "error": str(error)
        }


# ============================================================
# COMPLIANCE REVIEW
# ============================================================

def compliance_review(analysis):

    review_items = []

    promo_keywords = analysis.get(
        "promo_keywords",
        []
    )

    promo_slide_refs = analysis.get(
        "promo_slide_refs",
        []
    )

    review_keywords = analysis.get(
        "review_keywords",
        []
    )

    review_slide_refs = analysis.get(
        "review_slide_refs",
        []
    )

    it_keywords = analysis.get(
        "it_keywords",
        []
    )

    it_slide_refs = analysis.get(
        "it_slide_refs",
        []
    )

    if promo_keywords:

        review_items.append(
            {
                "category":
                    "Promotional Content",

                "severity":
                    "High",

                "status":
                    "Review Required",

                "message":
                    (
                        "Promotional-related keywords were detected. "
                        "The material requires further compliance review."
                    ),

                "keywords":
                    promo_keywords,

                "slide_refs":
                    promo_slide_refs,
            }
        )

    if review_keywords:

        review_items.append(
            {
                "category":
                    "Compliance-Sensitive Content",

                "severity":
                    "Review",

                "status":
                    "Review Required",

                "message":
                    (
                        "Terms that may require compliance interpretation "
                        "were detected. Keyword detection alone cannot "
                        "determine compliance."
                    ),

                "keywords":
                    review_keywords,

                "slide_refs":
                    review_slide_refs,
            }
        )

    if it_keywords:

        review_items.append(
            {
                "category":
                    "IT / Support Content",

                "severity":
                    "Info",

                "status":
                    "Informational",

                "message":
                    (
                        "Internal IT / support-related content was detected."
                    ),

                "keywords":
                    it_keywords,

                "slide_refs":
                    it_slide_refs,
            }
        )

    if not review_items:

        review_items.append(
            {
                "category":
                    "Initial Screening",

                "severity":
                    "Info",

                "status":
                    "No Obvious Issues",

                "message":
                    (
                        "No obvious promotional or compliance-sensitive "
                        "keywords were detected by the current screening rules."
                    ),

                "keywords":
                    [],

                "slide_refs":
                    [],
            }
        )

    overall_status = (
        "Review Required"
        if promo_keywords or review_keywords
        else "Pending AI Review"
    )

    return {
        "status": overall_status,
        "review_items": review_items,
    }


# ============================================================
# DISPLAY SLIDE REFERENCES
# ============================================================

def display_slide_references(slide_refs):

    if not slide_refs:
        return

    sorted_refs = sorted(
        slide_refs,
        key=lambda item: (
            int(item[0]),
            str(item[1]).lower()
        )
    )

    for slide_number, keyword in sorted_refs:

        st.write(
            f"• Slide {slide_number} — `{keyword}`"
        )


# ============================================================
# DISPLAY MATERIAL ANALYSIS
# ============================================================

def display_analysis(analysis):

    st.divider()

    st.header(
        "🤖 Material Analysis"
    )

    col1, col2, col3 = st.columns(3)

    with col1:
        st.metric(
            "Material Type",
            analysis["material_type"]
        )

    with col2:
        st.metric(
            "Risk Level",
            analysis["risk_level"]
        )

    with col3:
        st.metric(
            "Compliance Status",
            analysis["compliance_status"]
        )

    st.subheader("Findings")

    for finding in analysis["findings"]:
        st.write(
            f"• {finding}"
        )

    if analysis["it_keywords"]:

        st.subheader(
            "💻 Detected IT / Support Keywords"
        )

        st.write(
            ", ".join(
                analysis["it_keywords"]
            )
        )

    if analysis["promo_keywords"]:

        st.subheader(
            "⚠️ Detected Promotional Keywords"
        )

        st.write(
            ", ".join(
                analysis["promo_keywords"]
            )
        )

    if analysis["review_keywords"]:

        st.subheader(
            "🔍 Compliance-Sensitive Keywords"
        )

        st.write(
            ", ".join(
                analysis["review_keywords"]
            )
        )

    if analysis["it_slide_refs"]:

        st.subheader(
            "📌 IT / Support Slide References"
        )

        display_slide_references(
            analysis["it_slide_refs"]
        )

    if analysis["promo_slide_refs"]:

        st.subheader(
            "📌 Promotional Slide References"
        )

        display_slide_references(
            analysis["promo_slide_refs"]
        )

    if analysis["review_slide_refs"]:

        st.subheader(
            "📌 Compliance Review Slide References"
        )

        display_slide_references(
            analysis["review_slide_refs"]
        )


# ============================================================
# DISPLAY EVIDENCE / CONTEXT
# ============================================================

def display_evidence_context(evidence_items):

    st.divider()

    st.header(
        "📜 Evidence / Context"
    )

    st.caption(
        "Source context captured from the slide where the keyword was detected."
    )

    if not evidence_items:

        st.info(
            "No evidence was generated from the current screening."
        )

        return

    for item in evidence_items:

        st.write(
            f"• Slide {item['slide']} — "
            + ", ".join(
                f"`{keyword}`"
                for keyword in item["keywords"]
            )
        )

        st.caption(
            f"Category: {item['category']}"
        )

        st.write(
            f"Evidence: {item['evidence']}"
        )


# ============================================================
# DISPLAY AI CONTEXT ANALYSIS
# ============================================================

def display_ai_context_analysis(ai_result):

    st.divider()

    st.header(
        "🧠 AI Context Analysis"
    )

    status = ai_result.get(
        "status",
        "Unknown"
    )

    if status == "Completed":

        st.success(
            "AI contextual analysis completed."
        )

    elif status == "AI Not Configured":

        st.info(
            "AI Context Analysis is ready but not connected "
            "to an approved AI endpoint yet."
        )

    elif status == "AI Error":

        st.error(
            "AI Context Analysis encountered an error."
        )

    elif status == "No Evidence":

        st.info(
            "No evidence is available for AI analysis."
        )

    summary = ai_result.get(
        "overall_summary",
        ""
    )

    if summary:
        st.write(
            f"**Overall Summary:** {summary}"
        )

    items = ai_result.get(
        "items",
        []
    )

    if not items:

        if ai_result.get("error"):

            st.caption(
                f"Technical detail: {ai_result['error']}"
            )

        return

    items = sorted(
        items,
        key=lambda item: (
            int(item.get("slide", 0)),
            str(item.get("category", "")).lower()
        )
    )

    for item in items:

        st.subheader(
            f"Slide {item.get('slide', '?')} — "
            f"{item.get('assessment', 'Assessment')}"
        )

        st.write(
            f"**Category:** "
            f"{item.get('category', 'N/A')}"
        )

        keywords = item.get(
            "keywords",
            []
        )

        if keywords:

            st.write(
                "**Detected terms:** "
                + ", ".join(keywords)
            )

        st.write(
            f"**Context:** "
            f"{item.get('context', 'N/A')}"
        )

        st.write(
            f"**Risk Level:** "
            f"{item.get('risk_level', 'N/A')}"
        )

        st.write(
            f"**Reason:** "
            f"{item.get('reason', 'N/A')}"
        )

        st.write(
            f"**Recommendation:** "
            f"{item.get('recommendation', 'N/A')}"
        )


# ============================================================
# DISPLAY COMPLIANCE REVIEW
# ============================================================

def display_compliance_review(review):

    st.divider()

    st.header(
        "🛡️ Compliance Review"
    )

    if review["status"] == "Review Required":

        st.warning(
            "Review Required — the material contains "
            "content that needs further compliance interpretation."
        )

    else:

        st.info(
            "Pending AI Review — no obvious issue was identified "
            "by the current rule-based screening."
        )

    for item in review["review_items"]:

        st.subheader(
            f"{item['category']} — {item['status']}"
        )

        st.write(
            f"**Severity:** {item['severity']}"
        )

        st.write(
            item["message"]
        )

        if item["keywords"]:

            st.write(
                "**Detected terms:** "
                + ", ".join(
                    item["keywords"]
                )
            )

        if item["slide_refs"]:

            st.write(
                "**Slide references:**"
            )

            display_slide_references(
                item["slide_refs"]
            )

    st.info(
        "Prototype pre-review only. "
        "Rule-based screening and AI contextual analysis do not "
        "make a final compliance determination. "
        "Human review remains required for applicable content."
    )


# ============================================================
# MAIN APPLICATION
# ============================================================

st.title(
    "🔎 Material Scanner Agent"
)

st.write(
    APP_SUBTITLE
)

st.divider()


# ============================================================
# UPLOAD MATERIAL
# ============================================================

st.header(
    "Upload Material"
)

uploaded_file = st.file_uploader(
    "Upload your material",
    type=["pptx"],
    help="Currently supports PowerPoint PPTX files."
)


# ============================================================
# PROCESS UPLOADED FILE
# ============================================================

if uploaded_file:

    st.success(
        f"Uploaded: {uploaded_file.name}"
    )

    try:

        # ====================================================
        # POWERPOINT EXTRACTION
        # ====================================================

        slides_data = extract_slides(
            uploaded_file
        )

        st.info(
            f"Material received. "
            f"{len(slides_data)} slides detected."
        )

        # ====================================================
        # EXTRACTED CONTENT
        # ====================================================

        st.divider()

        st.header(
            "📄 Extracted Content"
        )

        for slide in slides_data:

            st.subheader(
                f"Slide {slide['slide']}"
            )

            if slide["texts"]:

                for text in slide["texts"]:

                    st.write(
                        text
                    )

            else:

                st.caption(
                    "No text content detected on this slide."
                )

        # ====================================================
        # RULE-BASED MATERIAL ANALYSIS
        # ====================================================

        analysis = analyze_material(
            slides_data
        )

        # ====================================================
        # EVIDENCE / CONTEXT EXTRACTION
        # ====================================================

        evidence_items = build_evidence_context(
            slides_data,
            analysis
        )

        analysis["evidence"] = evidence_items

        display_analysis(
            analysis
        )

        display_evidence_context(
            evidence_items
        )

        # ====================================================
        # AI CONTEXT ANALYSIS
        # ====================================================

        with st.spinner(
            "Running AI Context Analysis..."
        ):

            ai_result = ai_context_analysis(
                analysis,
                evidence_items
            )

        display_ai_context_analysis(
            ai_result
        )

        # ====================================================
        # COMPLIANCE REVIEW
        # ====================================================

        review = compliance_review(
            analysis
        )

        display_compliance_review(
            review
        )

    except Exception as error:

        st.error(
            "Unable to process this PowerPoint file."
        )

        st.exception(
            error
        )

else:

    st.info(
        "Please upload a PowerPoint PPTX file to begin."
    )
