from pptx import Presentation


def extract_pptx(file_path: str) -> list[dict]:
    """
    Extract readable text from each slide in a PowerPoint file.

    Returns:
        [
            {
                "slide_number": 1,
                "text": "..."
            },
            ...
        ]
    """

    presentation = Presentation(file_path)

    slides = []

    for slide_number, slide in enumerate(presentation.slides, start=1):
        text_parts = []

        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text.strip()

            if text:
                text_parts.append(text)

        slide_text = "\n".join(text_parts)

        slides.append(
            {
                "slide_number": slide_number,
                "text": slide_text,
            }
        )

    return slides


if __name__ == "__main__":
    print("PPTX Extractor module loaded.")